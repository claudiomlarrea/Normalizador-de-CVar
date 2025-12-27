import io
import re
import pandas as pd
import streamlit as st
from docx import Document

from scorer import load_criteria, score_text

st.set_page_config(page_title="Valorador de CVar CLEAN", layout="wide")
st.title("Valorador de CVar (TXT limpio)")
st.caption("Entrada: *_CVAR_CLEAN.txt (salida del Normalizador). Salidas: Excel + Word con puntajes y categoría.")

# Cargar criteria.json
try:
    criteria = load_criteria("criteria.json")
except Exception as e:
    st.error(f"No se pudo leer criteria.json: {e}")
    st.stop()

uploaded = st.file_uploader(
    "Subí el archivo TXT limpio (CVAR_CLEAN)",
    type=["txt"]
)

if not uploaded:
    st.info("Esperando archivo TXT limpio...")
    st.stop()

# Leer texto completo
raw = uploaded.read()
text = raw.decode("utf-8", errors="ignore")

# Vista previa (solo para UI)
preview_lines = 200
lines = text.splitlines()
preview = "\n".join(lines[:preview_lines])

st.subheader(f"Vista previa (primeras {preview_lines} líneas)")
st.code(preview if preview.strip() else "(archivo vacío)", language="text")

# Puntuar TODO el texto
with st.spinner("Calculando puntajes..."):
    results, section_totals, total_points, category, categorias = score_text(text, criteria)

# Resumen
c1, c2, c3 = st.columns([1, 1, 2])
c1.metric("Puntaje total", f"{total_points:.1f}")
c2.metric("Categoría", category)
desc = categorias.get(category, {}).get("descripcion", "")
c3.write(desc)

# DataFrame detalle
rows = []
for r in results:
    rows.append({
        "Sección": r.section,
        "Ítem": r.item,
        "Conteo": r.count,
        "Unit points": r.unit_points,
        "Puntos bruto": r.raw_points,
        "Tope ítem": r.item_max_points,
        "Puntos (tope aplicado)": r.capped_item_points,
        "Evidencia (1er match)": r.evidence
    })
df = pd.DataFrame(rows)

st.subheader("Detalle por ítem")
st.dataframe(df, use_container_width=True, hide_index=True)

st.subheader("Totales por sección (tope de sección aplicado)")
df_sec = pd.DataFrame([{"Sección": k, "Puntos": v} for k, v in section_totals.items()]).sort_values("Puntos", ascending=False)
st.dataframe(df_sec, use_container_width=True, hide_index=True)

# ---- Export Excel ----
excel_buf = io.BytesIO()
with pd.ExcelWriter(excel_buf, engine="xlsxwriter") as writer:
    df.to_excel(writer, sheet_name="Detalle", index=False)
    df_sec.to_excel(writer, sheet_name="Secciones", index=False)
    pd.DataFrame([{
        "Archivo": uploaded.name,
        "Puntaje total": total_points,
        "Categoría": category,
        "Descripción": desc
    }]).to_excel(writer, sheet_name="Resumen", index=False)
excel_buf.seek(0)

# ---- Export Word ----
doc = Document()
doc.add_heading("Informe de valoración — CVar (TXT limpio)", level=1)
doc.add_paragraph(f"Archivo evaluado: {uploaded.name}")
doc.add_paragraph(f"Puntaje total: {total_points:.1f}")
doc.add_paragraph(f"Categoría: {category}")
if desc:
    doc.add_paragraph(desc)

doc.add_heading("Totales por sección", level=2)
for _, row in df_sec.iterrows():
    doc.add_paragraph(f"- {row['Sección']}: {float(row['Puntos']):.1f}")

doc.add_heading("Ítems detectados (conteo > 0) con evidencia", level=2)
for _, row in df.iterrows():
    if int(row["Conteo"]) > 0:
        doc.add_paragraph(
            f"• {row['Sección']} — {row['Ítem']}: {float(row['Puntos (tope aplicado)']):.1f} pts (conteo={int(row['Conteo'])})"
        )
        ev = str(row.get("Evidencia (1er match)", "")).strip()
        if ev:
            doc.add_paragraph(f"Evidencia: {ev}")

word_buf = io.BytesIO()
doc.save(word_buf)
word_buf.seek(0)

st.download_button(
    "⬇️ Descargar Excel (puntajes)",
    data=excel_buf,
    file_name=uploaded.name.replace(".txt", "") + "__PUNTAJE.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True
)

st.download_button(
    "⬇️ Descargar Word (informe)",
    data=word_buf,
    file_name=uploaded.name.replace(".txt", "") + "__INFORME.docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    use_container_width=True
)
